"""Korpus yukleyici.

Excel'deki 3 korpus sayfasindan metinleri okur ve Olcum_Sablonu'ndaki
metin_id'lerle eslestirir.

DIKKAT — ID tutarsizligi:
    G2 sayfasi:  satir 1-30 -> 'G2_01'..'G2_30'
                 satir 31-100 -> 'G2_YZ_31'..'G2_YZ_100'
    G3 sayfasi:  satir 1-30 -> 'G3_01'..'G3_30'
                 satir 31-100 -> 'G3_Insanlastirilmis_31'..
    Olcum_Sablonu ise hep uzun formu kullanir ('G2_YZ_01').

Bu yuzden eslestirme ID string'iyle YAPILAMAZ. ID'nin sonundaki sayi her uc
sayfada da 1..100 sirasiyla gittigi dogrulandi; eslestirme o sayi uzerinden
yapilir.
"""
from __future__ import annotations

import re
from dataclasses import dataclass

import openpyxl

from . import config

# (sayfa adi, metin sutunu, Olcum_Sablonu'ndaki grup adi)
KAYNAKLAR = [
    ("G1_Insan_Korpus", 9, "G1_Insan"),
    ("G2_YZ_Korpus", 5, "G2_YZ"),
    ("G3_Insanlastirilmis", 3, "G3_Insanlastirilmis"),
]


@dataclass(frozen=True)
class Metin:
    metin_id: str      # Olcum_Sablonu formatinda, orn. 'G2_YZ_01'
    grup: str          # 'G1_Insan' | 'G2_YZ' | 'G3_Insanlastirilmis'
    gercek_etiket: str # 'Insan' | 'YZ'
    dil: str           # 'TUR' | 'ENG'
    metin: str

    @property
    def kelime(self) -> int:
        return len(self.metin.split())


def _son_sayi(deger) -> int:
    """ID'nin sonundaki sayiyi cikarir: 'G2_YZ_31' -> 31, 'G2_01' -> 1."""
    bulunan = re.findall(r"(\d+)", str(deger))
    if not bulunan:
        raise ValueError(f"ID'de sayi yok: {deger!r}")
    return int(bulunan[-1])


def yukle(xlsx=None) -> dict[str, Metin]:
    """Tum korpusu {metin_id: Metin} olarak dondurur."""
    wb = openpyxl.load_workbook(xlsx or config.XLSX, data_only=True)

    # Olcum_Sablonu tek dogru kaynak: metin_id, grup, etiket, dil buradan gelir
    olc = wb["Olcum_Sablonu"]
    kimlik: dict[tuple[str, int], tuple[str, str, str]] = {}
    for r in range(2, olc.max_row + 1):
        mid = olc.cell(r, 1).value
        if not mid:
            continue
        grup = olc.cell(r, 2).value
        anahtar = (grup, _son_sayi(mid))
        if anahtar not in kimlik:
            kimlik[anahtar] = (mid, olc.cell(r, 3).value, olc.cell(r, 4).value)

    sonuc: dict[str, Metin] = {}
    for sayfa, sutun, grup in KAYNAKLAR:
        ws = wb[sayfa]
        for r in range(2, ws.max_row + 1):
            ham_id = ws.cell(r, 1).value
            if ham_id is None:
                continue
            no = _son_sayi(ham_id)
            anahtar = (grup, no)
            if anahtar not in kimlik:
                raise KeyError(
                    f"{sayfa} satir {r} (id={ham_id}) Olcum_Sablonu'nda karsiliksiz"
                )
            mid, etiket, dil = kimlik[anahtar]
            metin = ws.cell(r, sutun).value
            if not metin or not str(metin).strip():
                raise ValueError(f"{sayfa} satir {r}: metin bos")
            sonuc[mid] = Metin(
                metin_id=mid,
                grup=grup,
                gercek_etiket=etiket,
                dil=dil,
                metin=str(metin).strip(),
            )

    beklenen = 300
    if len(sonuc) != beklenen:
        raise AssertionError(f"{beklenen} metin bekleniyordu, {len(sonuc)} yuklendi")
    return sonuc


if __name__ == "__main__":
    m = yukle()
    print(f"{len(m)} metin yuklendi")
    for grup in ("G1_Insan", "G2_YZ", "G3_Insanlastirilmis"):
        alt = [x for x in m.values() if x.grup == grup]
        tur = sum(1 for x in alt if x.dil == "TUR")
        print(
            f"  {grup:22s} n={len(alt):3d} | TUR={tur:3d} ENG={len(alt)-tur:3d} "
            f"| ort kelime={sum(x.kelime for x in alt)/len(alt):.1f}"
        )
