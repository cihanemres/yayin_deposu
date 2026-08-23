"""JSONL sonuclarini Excel'e yazar.

Ayri adim olmasinin sebebi: olcum saatler surebilir ve Excel acikken yazilamaz.
Once tum veri JSONL'e toplanir, sonra tek seferde Excel'e islenir.
"""
from __future__ import annotations

import json
import shutil
from datetime import datetime

import openpyxl

from . import config

ARAC_SIRASI = ["GPTZero", "ZeroGPT", "Copyleaks", "Sapling", "Originality_ai", "Pangram"]


def yedekle() -> str:
    damga = datetime.now().strftime("%Y%m%d_%H%M")
    hedef = config.XLSX.with_name(f"{config.XLSX.stem}_YEDEK_{damga}.xlsx")
    shutil.copy2(config.XLSX, hedef)
    return str(hedef)


def _kilitli() -> bool:
    return (config.XLSX.parent / f"~${config.XLSX.name}").exists()


def _sonuclari_oku() -> dict[tuple[str, str], dict]:
    """En son basarili olcumu (metin_id, arac) basina dondurur."""
    son: dict[tuple[str, str], dict] = {}
    if not config.RESULTS_JSONL.exists():
        return son
    with config.RESULTS_JSONL.open(encoding="utf-8") as f:
        for satir in f:
            satir = satir.strip()
            if not satir:
                continue
            try:
                k = json.loads(satir)
            except json.JSONDecodeError:
                continue
            if k.get("yz_olasilik") is None:
                continue
            son[(k["metin_id"], k["arac"])] = k  # sonraki eskisini ezer
    return son


def pangram_satirlari_ekle(ws) -> int:
    """Her metin_id icin eksik Pangram satirini ekler (1500 -> 1800)."""
    mevcut = {(ws.cell(r, 1).value, ws.cell(r, 5).value)
              for r in range(2, ws.max_row + 1)}
    # metin_id -> (grup, etiket, dil)
    meta: dict[str, tuple] = {}
    for r in range(2, ws.max_row + 1):
        mid = ws.cell(r, 1).value
        if mid and mid not in meta:
            meta[mid] = (ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value)

    eklenen = 0
    satir = ws.max_row + 1
    for mid, (grup, etiket, dil) in meta.items():
        if (mid, "Pangram") in mevcut:
            continue
        ws.cell(satir, 1, mid)
        ws.cell(satir, 2, grup)
        ws.cell(satir, 3, etiket)
        ws.cell(satir, 4, dil)
        ws.cell(satir, 5, "Pangram")
        satir += 1
        eklenen += 1
    return eklenen


def yaz(pangram_ekle: bool = True, gruplar: list[str] | None = None) -> None:
    if _kilitli():
        raise RuntimeError(
            f"Excel dosyasi acik gorunuyor (~${config.XLSX.name} lock dosyasi var).\n"
            "Excel'i kapatip tekrar deneyin."
        )

    sonuclar = _sonuclari_oku()
    if not sonuclar:
        print("! JSONL'de basarili sonuc yok. Once `python run.py --olc` calistirin.")
        return

    grup_kumesi = set(gruplar) if gruplar else None
    if grup_kumesi:
        print(f"Grup filtresi: {', '.join(sorted(grup_kumesi))} "
              "(digerleri el degmeden birakilir)")

    yedek = yedekle()
    print(f"Yedek alindi: {yedek}")

    wb = openpyxl.load_workbook(config.XLSX)
    ws = wb["Olcum_Sablonu"]

    # Grup filtresi varken Pangram satiri ekleme adimini atla: filtrelenen
    # gruplara mudahale etmemek icin (satirlar zaten sablonda mevcut).
    if pangram_ekle and not grup_kumesi:
        n = pangram_satirlari_ekle(ws)
        if n:
            print(f"Pangram satiri eklendi: {n} (toplam {ws.max_row - 1} olcum satiri)")

    yazilan = bos = atlanan = 0
    for r in range(2, ws.max_row + 1):
        mid, arac = ws.cell(r, 1).value, ws.cell(r, 5).value
        if grup_kumesi is not None and ws.cell(r, 2).value not in grup_kumesi:
            atlanan += 1
            continue
        k = sonuclar.get((mid, arac))
        if not k:
            bos += 1
            continue
        ws.cell(r, 6, k["yz_olasilik"])                      # yz_olasilik_%
        ws.cell(r, 7, k["karar"])                            # karar_YZ_Insan
        ws.cell(r, 8, datetime.fromisoformat(k["tarih"]))    # tarih
        ws.cell(r, 9, k.get("etiket") or None)               # not
        yazilan += 1

    # Yontem_Notu DURUM guncelle (grup filtresi varken atla — sayim yaniltici olur)
    yn = wb["Yontem_Notu"] if not grup_kumesi else None
    for r in range(1, (yn.max_row if yn else 0) + 1):
        if yn.cell(r, 1).value == "DURUM":
            eski = str(yn.cell(r, 2).value or "")
            temel = eski.split("| Ölçüm:")[0].strip()
            yn.cell(r, 2, f"{temel} | Ölçüm: {yazilan}/{ws.max_row - 1} (API)")
        if yn.cell(r, 1).value == "Araçlar":
            yn.cell(r, 2, "GPTZero, ZeroGPT, Copyleaks, Sapling, Originality.ai, "
                          "Pangram — tümü API üzerinden, tek protokol")

    wb.save(config.XLSX)
    ek = f" | Filtre disi (atlanan): {atlanan}" if grup_kumesi else ""
    print(f"Yazildi: {yazilan} satir | Bos kalan: {bos}{ek}")
    print(f"Kaydedildi: {config.XLSX}")
