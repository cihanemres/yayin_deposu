"""G3 Islem_Notu'nun iki bicimi (TUR: kisaltma yasagi yazili / ENG: yazili degil)
arasinda UYGULAMA farki var mi? Uzunluk orani uzerinden sinanir."""
import openpyxl, statistics as st
from scipy import stats

wb = openpyxl.load_workbook(
    r"C:\Users\cihanemres\Desktop\icits\YZ_Tespit_Korpus_v9.xlsx",
    read_only=True, data_only=True)

g2 = {}
for r in wb["G2_YZ_Korpus"].iter_rows(min_row=2, values_only=True):
    if r[0]:
        g2[str(r[0])] = {"kelime": r[5], "metin": r[4] or ""}

g3 = []
for r in wb["G3_Insanlastirilmis"].iter_rows(min_row=2, values_only=True):
    if not r[0]:
        continue
    g3.append({"id": str(r[0]), "kaynak": str(r[1]), "metin": r[2] or "",
               "kelime": r[3], "not": (r[4] or "").strip()})

# not bicimlerini ayir
notlar = sorted({x["not"] for x in g3})
kisaltma_yasagi = [n for n in notlar if "kısaltma" in n.lower()]
assert len(kisaltma_yasagi) == 1, notlar
KY = kisaltma_yasagi[0]

def dil(mid):
    n = int(mid.rsplit("_", 1)[1])
    return "TUR" if n <= 50 else "ENG"

satir = []
for x in g3:
    # sozcuk sayilarini metinden hesapla (whitespace token) — Kelime hucresi bos olabiliyor
    w3 = len(x["metin"].split())
    w2 = len(g2[x["kaynak"]]["metin"].split())
    satir.append({
        "id": x["id"], "dil": dil(x["id"]),
        "not_ky": x["not"] == KY,
        "oran_metin": w3 / w2,
        "w2": w2, "w3": w3,
    })

print("=== not bicimi x dil capraz tablosu ===")
for d in ("TUR", "ENG"):
    ky = sum(1 for s in satir if s["dil"] == d and s["not_ky"])
    print("  %s: kisaltma-yasagi yazili %2d / %2d" % (d, ky, sum(1 for s in satir if s["dil"] == d)))

print("\n=== uzunluk orani (G3/G2), metinden hesaplanan ===")
gruplar = {}
for s in satir:
    gruplar.setdefault(s["not_ky"], []).append(s["oran_metin"])

for ky in (True, False):
    v = gruplar[ky]
    etiket = "NOT 1 (kisaltma yasagi YAZILI)" if ky else "NOT 2 (yazili DEGIL)"
    print("  %-34s n=%d  ort=%.4f  ss=%.4f  min=%.4f  max=%.4f"
          % (etiket, len(v), st.mean(v), st.stdev(v), min(v), max(v)))

a, b = gruplar[True], gruplar[False]
t, p = stats.ttest_ind(a, b, equal_var=False)
u, pu = stats.mannwhitneyu(a, b)
d = (st.mean(a) - st.mean(b)) / (((st.stdev(a)**2 + st.stdev(b)**2) / 2) ** 0.5)
print("\n  Welch t = %.4f, p = %.4f | Mann-Whitney U = %.1f, p = %.4f | Cohen d = %.3f"
      % (t, p, u, pu, d))

# TOST esdegerlik: fark +-0.05 bandinda mi
from math import sqrt
delta = 0.05
se = sqrt(st.stdev(a)**2/len(a) + st.stdev(b)**2/len(b))
diff = st.mean(a) - st.mean(b)
df = len(a) + len(b) - 2
t1 = (diff + delta) / se
t2 = (diff - delta) / se
p1 = 1 - stats.t.cdf(t1, df)
p2 = stats.t.cdf(t2, df)
print("  TOST (delta=+-0.05): fark=%.4f, p_alt=%.4g, p_ust=%.4g -> esdegerlik %s"
      % (diff, p1, p2, "KANITLANDI" if max(p1, p2) < .05 else "kanitlanamadi"))

# bant ihlalleri
print("\n=== +-%10 bandi ihlalleri ===")
for ky in (True, False):
    ihl = [s for s in satir if s["not_ky"] == ky and abs(s["oran_metin"] - 1) > 0.10]
    etiket = "NOT 1" if ky else "NOT 2"
    print("  %s: %d ihlal  %s" % (etiket, len(ihl), [(s["id"].replace("G3_Insanlastirilmis_", "G3_"), round(s["oran_metin"], 4)) for s in ihl]))
